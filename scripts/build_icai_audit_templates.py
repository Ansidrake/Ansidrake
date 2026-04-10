#!/usr/bin/env python3
"""Generate ICAI Standards on Auditing–aligned audit working paper templates (Word + Excel)."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "templates"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.space_after = Pt(6)


def add_label_value(doc: Document, label: str, value: str = "[To be completed]") -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{label}: ")
    r.bold = True
    p.add_run(value)


def thin_border() -> Border:
    s = Side(style="thin", color="000000")
    return Border(left=s, right=s, top=s, bottom=s)


def header_fill() -> PatternFill:
    return PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")


def style_header_row(ws, row: int, max_col: int) -> None:
    fill = header_fill()
    font = Font(bold=True, color="FFFFFF", size=11)
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border()


def build_word() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "ICAI-Standards-on-Auditing-Document-Templates.docx"
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Audit Documentation Templates\n(Aligned with ICAI Standards on Auditing)")
    r.bold = True
    r.font.size = Pt(16)
    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "These are generic placeholders. Tailor each section to the engagement, applicable SAs, "
        "firm methodology, and regulatory requirements. They do not replace professional judgment "
        "or the text of the Standards."
    )
    doc.add_page_break()

    sections = [
        (
            "1. Engagement acceptance / continuance and terms of audit",
            [
                ("Client name", ""),
                ("Financial year / period end", ""),
                ("Preconditions for an audit (SA 210)", "Document assessment."),
                ("Agreeing terms — audit engagement letter (SA 210)", "Reference / attach signed letter."),
                ("Independence and conflicts", "Document conclusions."),
                ("Resources and competence", ""),
            ],
        ),
        (
            "2. Overall audit strategy and audit plan (SA 300)",
            [
                ("Characteristics of the entity and environment", ""),
                ("Reporting objectives and applicable reporting framework", "Ind AS / AS / other."),
                ("Significant factors directing overall resource allocation", ""),
                ("Planned approach to materiality (SA 320)", "See materiality working paper."),
                ("Planned risk assessment procedures", ""),
                ("Planned further audit procedures — nature, timing, extent", ""),
                ("Direction, supervision, and review (SA 220)", ""),
            ],
        ),
        (
            "3. Understanding the entity and risk assessment (SA 315)",
            [
                ("Industry, regulatory, and other external factors", ""),
                ("Nature of the entity — operations, ownership, governance, business model", ""),
                ("Accounting policies and estimates", ""),
                ("Objectives, strategies, and related business risks", ""),
                ("Measurement and review of financial performance", ""),
                ("IT environment and general IT controls (as applicable)", ""),
                ("Identified risks of material misstatement — financial statement level", ""),
                ("Identified risks of material misstatement — assertion level", ""),
                ("Controls relevant to the audit — design and implementation", ""),
            ],
        ),
        (
            "4. Materiality and performance materiality (SA 320)",
            [
                ("Benchmark selected and rationale", ""),
                ("Percentage applied and rationale", ""),
                ("Overall materiality for financial statements as a whole", ""),
                ("Performance materiality", ""),
                ("Clearly trivial threshold (if used)", ""),
                ("Revisions during the audit", ""),
            ],
        ),
        (
            "5. Audit evidence — further procedures (SA 330, SA 500 series)",
            [
                ("Tests of controls — objective, population, sample, results", ""),
                ("Substantive analytical procedures — expectation, data, investigation", ""),
                ("Tests of details — objective, population, sample, results", ""),
                ("External confirmations (SA 505) — control log / exceptions", ""),
                ("Related parties (SA 550)", ""),
                ("Accounting estimates (SA 540)", ""),
                ("Going concern (SA 570)", ""),
                ("Subsequent events (SA 560)", ""),
            ],
        ),
        (
            "6. Evaluation of misstatements (SA 450)",
            [
                ("Accumulation of identified misstatements", ""),
                ("Uncorrected misstatements — nature, amount, classification", ""),
                ("Qualitative factors", ""),
                ("Communication with those charged with governance", ""),
                ("Written representations requested (SA 580)", ""),
            ],
        ),
        (
            "7. Completion and reporting (SA 700 series, SA 705, SA 706)",
            [
                ("Going concern — conclusions and disclosures", ""),
                ("Subsequent events — procedures and conclusions", ""),
                ("Written representations — obtained and evaluated", ""),
                ("Overall review of financial statements", ""),
                ("Form and content of auditor’s report", ""),
                ("Key audit matters / emphasis of matter / other matter (as applicable)", ""),
            ],
        ),
        (
            "8. Quality management at engagement level (SA 220, ISQM 1 engagement aspects)",
            [
                ("Engagement partner responsibilities", ""),
                ("Direction, supervision, and review — evidence on file", ""),
                ("Consultation / hot review (if applicable)", ""),
                ("Engagement quality review (if applicable)", ""),
            ],
        ),
    ]

    for idx, (heading, items) in enumerate(sections):
        add_heading(doc, heading, level=1)
        add_label_value(doc, "Engagement reference", "[Engagement no. / code]")
        add_label_value(doc, "Prepared by", "[Name / initials]")
        add_label_value(doc, "Date", "[DD/MM/YYYY]")
        doc.add_paragraph()
        doc.add_paragraph("Working paper content / conclusions:", style="List Bullet")
        for label, hint in items:
            line = f"{label}"
            if hint:
                line += f" — {hint}"
            doc.add_paragraph(line, style="List Bullet")
        doc.add_paragraph()
        doc.add_paragraph(
            "Prepared by: _________________    Reviewed by: _________________    Date: _________"
        )
        if idx < len(sections) - 1:
            doc.add_page_break()

    doc.save(path)
    return path


def add_sheet_table(
    wb: Workbook,
    name: str,
    headers: list[str],
    row_hints: list[str] | None = None,
    num_data_rows: int = 12,
) -> None:
    ws = wb.create_sheet(title=name[:31])
    for i, h in enumerate(headers, start=1):
        ws.cell(row=1, column=i, value=h)
    style_header_row(ws, 1, len(headers))
    start = 2
    if row_hints:
        for j, hint in enumerate(row_hints, start=start):
            ws.cell(row=j, column=1, value=hint)
        start = start + len(row_hints)
    for r in range(start, start + num_data_rows):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c, value="")
            ws.cell(row=r, column=c).border = thin_border()
    for col in range(1, len(headers) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(28, 12 + 2 * len(headers[col - 1]) // 3)


def build_excel() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / "ICAI-Standards-on-Auditing-Document-Templates.xlsx"
    wb = Workbook()
    default = wb.active
    default.title = "Index"
    default["A1"] = "ICAI Standards on Auditing — Audit documentation templates (Excel)"
    default["A1"].font = Font(bold=True, size=14)
    default["A3"] = "Instructions"
    default["A4"] = (
        "Use each sheet as a structured working paper. Link or cross-reference to detailed Word "
        "schedules where needed. Replace bracketed text with engagement-specific information."
    )
    default["A3"].font = Font(bold=True)
    rows = [
        ("Sheet", "Typical use (SA reference — illustrative)"),
        ("Engagement_Acceptance", "SA 210 — preconditions, engagement terms"),
        ("Audit_Strategy_Plan", "SA 300 — overall strategy and plan"),
        ("Risk_Assessment", "SA 315 — understanding and identified risks"),
        ("Materiality", "SA 320 — benchmarks, thresholds, revisions"),
        ("Further_Procedures_Log", "SA 330 — linking risks to procedures"),
        ("Misstatements_SA450", "SA 450 — accumulation and evaluation"),
        ("Completion_Checklist", "SA 500 series, SA 570, SA 560, reporting"),
    ]
    r0 = 6
    for i, (a, b) in enumerate(rows):
        default.cell(row=r0 + i, column=1, value=a)
        default.cell(row=r0 + i, column=2, value=b)
    default.column_dimensions["A"].width = 28
    default.column_dimensions["B"].width = 70

    add_sheet_table(
        wb,
        "Engagement_Acceptance",
        ["Item", "Description / procedure", "Evidence / reference", "Conclusion (Y/N/NA)", "Initials / date"],
        [
            "Preconditions documented",
            "Engagement letter signed / on file",
            "Independence confirmed",
            "Competence and resources",
        ],
    )
    add_sheet_table(
        wb,
        "Audit_Strategy_Plan",
        ["Section", "Key points", "Owner", "Status", "WP ref"],
        ["Reporting framework", "Significant components / locations", "Use of internal audit / experts", "IT / fraud considerations"],
    )
    add_sheet_table(
        wb,
        "Risk_Assessment",
        ["FS area / assertion", "Risks (ROMM)", "Controls relied upon?", "FAP planned (TOD/TOP/SAP)", "WP ref"],
        None,
        15,
    )
    add_sheet_table(
        wb,
        "Materiality",
        ["Element", "Amount / %", "Rationale", "Approved by", "Date"],
        ["Benchmark", "Overall materiality", "Performance materiality", "Clearly trivial (if any)", "Revisions"],
    )
    add_sheet_table(
        wb,
        "Further_Procedures_Log",
        ["Assertion", "Procedure ref.", "Nature / timing / extent", "Sample / population", "Result", "Exception?"],
        None,
        20,
    )
    add_sheet_table(
        wb,
        "Misstatements_SA450",
        ["Ref", "Description", "Debit / Credit", "Impact (P&L / BS)", "Corrected?", "TCTG communicated?"],
        None,
        15,
    )
    add_sheet_table(
        wb,
        "Completion_Checklist",
        ["Completion item", "Done (Y/N)", "WP ref / comment", "Reviewer"],
        [
            "Analytical review — overall FS",
            "Going concern (SA 570)",
            "Subsequent events (SA 560)",
            "Related parties (SA 550)",
            "Litigation & claims",
            "Written representations (SA 580)",
            "Audit report drafting / consistency check",
        ],
    )

    wb.save(path)
    return path


def main() -> None:
    w = build_word()
    x = build_excel()
    print(f"Written: {w}\nWritten: {x}")


if __name__ == "__main__":
    main()
